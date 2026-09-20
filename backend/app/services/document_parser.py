import os
import re
import io
import csv
from uuid import uuid4
from datetime import datetime, timezone
from typing import Any
import pypdf
import docx

from backend.app.models.documents import DocumentStatus, DocumentChunk

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv"}

DISALLOWED_EXTENSIONS = {
    ".exe", ".bat", ".cmd", ".sh", ".py", ".js", ".vbs", ".ps1", ".php",
    ".pl", ".cgi", ".dll", ".so", ".dylib", ".msi", ".jar", ".com", ".scr"
}

ALLOWED_MIME_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "text/plain",
    "text/markdown",
    "text/csv",
    "application/csv",
    "application/octet-stream",  # Fallback if browser sends default
}

PROMPT_INJECTION_PATTERNS = [
    r"ignore\s+all\s+previous\s+instructions",
    r"ignore\s+above\s+instructions",
    r"you\s+are\s+now\s+an\s+administrator",
    r"reveal\s+the\s+gemini\s+api\s+key",
    r"reveal\s+api\s+key",
    r"system\s*:",
    r"override\s+security",
]


def sanitize_untrusted_text(text: str) -> str:
    """Sanitize raw document text to neutralize prompt injection attacks."""
    sanitized = text
    for pattern in PROMPT_INJECTION_PATTERNS:
        sanitized = re.sub(pattern, "[REDACTED_PROMPT_INJECTION_ATTEMPT]", sanitized, flags=re.IGNORECASE)
    return sanitized


def sanitize_filename(filename: str) -> str:
    """Sanitize filename to prevent path traversal and unsafe filesystem characters."""
    filename = os.path.basename(filename)
    filename = re.sub(r"\.\.+", "", filename)  # strip ..
    filename = re.sub(r"[^\w\.\-]", "_", filename)
    filename = filename.lstrip(".")
    return filename or f"file_{uuid4().hex[:6]}"


def validate_upload_file(filename: str, content_bytes: bytes, content_type: str | None = None) -> None:
    """Validate file extension, executable rejection, size, and content non-emptiness."""
    if not content_bytes or len(content_bytes) == 0:
        raise ValueError("Uploaded file is empty.")

    if len(content_bytes) > MAX_FILE_SIZE:
        raise ValueError(f"File size ({round(len(content_bytes) / (1024 * 1024), 2)} MB) exceeds 10 MB limit.")

    ext = os.path.splitext(filename.lower())[1]
    if ext in DISALLOWED_EXTENSIONS or ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported or dangerous file format '{ext}'. Allowed formats: PDF, DOCX, TXT, MD, CSV.")


def extract_text_by_pages(filename: str, content_bytes: bytes) -> tuple[list[tuple[int, str]], str]:
    """Extract page-indexed text from supported document types.
    
    Returns:
        tuple[list[tuple[page_number, page_text]], status_string]
    """
    ext = os.path.splitext(filename.lower())[1]
    pages: list[tuple[int, str]] = []

    if ext == ".pdf":
        try:
            reader = pypdf.PdfReader(io.BytesIO(content_bytes))
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append((idx + 1, text.strip()))
        except Exception as e:
            raise ValueError(f"Failed to extract PDF content: {e}")

        if not pages:
            return [], DocumentStatus.NO_EXTRACTABLE_TEXT.value

    elif ext == ".docx":
        try:
            doc = docx.Document(io.BytesIO(content_bytes))
            lines = []
            for p in doc.paragraphs:
                if p.text.strip():
                    lines.append(p.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_txt:
                        lines.append(row_txt)
            full_text = "\n\n".join(lines)
            if full_text.strip():
                pages.append((1, full_text.strip()))
        except Exception as e:
            raise ValueError(f"Failed to extract DOCX content: {e}")

    elif ext in (".txt", ".md"):
        try:
            text = content_bytes.decode("utf-8", errors="replace").strip()
            if text:
                pages.append((1, text))
        except Exception as e:
            raise ValueError(f"Failed to extract text content: {e}")

    elif ext == ".csv":
        try:
            text_content = content_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(text_content))
            rows = list(reader)
            if rows:
                header = rows[0]
                lines = []
                for row in rows[1:]:
                    if not any(row):
                        continue
                    if len(header) == len(row):
                        pair_str = ", ".join(f"{h.strip()}: {v.strip()}" for h, v in zip(header, row) if v.strip())
                    else:
                        pair_str = ", ".join(v.strip() for v in row if v.strip())
                    if pair_str:
                        lines.append(pair_str)
                full_csv_txt = "\n".join(lines)
                if full_csv_txt.strip():
                    pages.append((1, full_csv_txt.strip()))
        except Exception as e:
            raise ValueError(f"Failed to extract CSV content: {e}")

    if not pages:
        return [], DocumentStatus.NO_EXTRACTABLE_TEXT.value

    return pages, DocumentStatus.READY.value


def chunk_document_pages(
    document_id: str,
    user_id: str,
    filename: str,
    pages: list[tuple[int, str]],
    target_chunk_size: int = 800,
    overlap: int = 100,
) -> list[DocumentChunk]:
    """Chunk extracted page texts into deterministic paragraph-aware chunks."""
    chunks: list[DocumentChunk] = []
    global_chunk_idx = 0

    for page_num, page_text in pages:
        paragraphs = [p.strip() for p in page_text.split("\n\n") if p.strip()]
        if not paragraphs:
            paragraphs = [page_text.strip()]

        current_chunk_paragraphs = []
        current_len = 0

        for p in paragraphs:
            p_len = len(p)
            if current_len + p_len > target_chunk_size and current_chunk_paragraphs:
                chunk_text = "\n\n".join(current_chunk_paragraphs)
                chunks.append(
                    DocumentChunk(
                        chunk_id=f"{document_id}-chunk-{global_chunk_idx}",
                        document_id=document_id,
                        user_id=user_id,
                        text=chunk_text,
                        page_number=page_num,
                        chunk_index=global_chunk_idx,
                        metadata={
                            "filename": filename,
                            "page_number": page_num,
                            "chunk_index": global_chunk_idx,
                        },
                    )
                )
                global_chunk_idx += 1

                # Retain overlap paragraph if possible
                overlap_text = current_chunk_paragraphs[-1] if len(current_chunk_paragraphs) > 1 else ""
                current_chunk_paragraphs = [overlap_text, p] if overlap_text else [p]
                current_len = sum(len(x) for x in current_chunk_paragraphs)
            else:
                current_chunk_paragraphs.append(p)
                current_len += p_len

        if current_chunk_paragraphs:
            chunk_text = "\n\n".join(current_chunk_paragraphs)
            chunks.append(
                DocumentChunk(
                    chunk_id=f"{document_id}-chunk-{global_chunk_idx}",
                    document_id=document_id,
                    user_id=user_id,
                    text=chunk_text,
                    page_number=page_num,
                    chunk_index=global_chunk_idx,
                    metadata={
                        "filename": filename,
                        "page_number": page_num,
                        "chunk_index": global_chunk_idx,
                    },
                )
            )
            global_chunk_idx += 1

    return chunks
